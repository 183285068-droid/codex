"""Check installed versions and exercise the scientific Python environment."""
import argparse
import importlib.metadata
import json
import platform
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--versions-only', action='store_true')
args = parser.parse_args()
versions = {}
errors = []
for requirement in (ROOT / 'requirements.txt').read_text().splitlines():
    name, expected = requirement.split('==')
    try:
        actual = importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        actual = 'missing'
    versions[name] = actual
    if actual != expected:
        errors.append(f'{name}: expected {expected}, found {actual}')
if errors:
    print('\n'.join(errors), file=sys.stderr)
    sys.exit(1)
if args.versions_only:
    print('All pinned packages are available.')
    sys.exit(0)

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.linear_model import LinearRegression
from openpyxl import load_workbook

out = ROOT / 'outputs' / 'environment-check'
out.mkdir(parents=True, exist_ok=True)
frame = pd.DataFrame({'x': [1., 2., 3., 4., 5.], 'y': [3., 5., 7., 9., 11.]})
frame.to_csv(out / 'example.csv', index=False)
pd.testing.assert_frame_equal(frame, pd.read_csv(out / 'example.csv'))
frame.to_excel(out / 'example.xlsx', index=False)
pd.testing.assert_frame_equal(frame, pd.read_excel(out / 'example.xlsx'), check_dtype=False)
book = load_workbook(out / 'example.xlsx', read_only=True)
assert book.active.max_row == 6
book.close()
model = LinearRegression().fit(frame[['x']], frame['y'])
assert np.isclose(model.coef_[0], 2) and np.isclose(model.intercept_, 1)
result = stats.ttest_ind([1, 2, 3, 4, 5], [2, 3, 4, 5, 6], equal_var=False)
assert np.isclose(result.statistic, -1) and np.isfinite(result.pvalue)
fig, ax = plt.subplots(figsize=(5, 3))
sns.scatterplot(data=frame, x='x', y='y', ax=ax)
ax.set_title('Environment check: y = 2x + 1')
fig.tight_layout()
fig.savefig(out / 'example.png', dpi=120)
plt.close(fig)
report = {'status': 'PASS', 'python': platform.python_version(), 'packages': versions,
          'checks': ['CSV read/write', 'Excel read/write', 'Welch t-test',
                     'linear regression', 'headless PNG plotting'],
          'welch_t': float(result.statistic), 'welch_p': float(result.pvalue),
          'limitations': ['R and Rscript are not installed.',
                          'Cloud workspace is temporary; use setup.sh to restore.',
                          'External network access is governed by the host.']}
(out / 'report.json').write_text(json.dumps(report, indent=2, ensure_ascii=False) + '\n')
print(json.dumps(report, indent=2, ensure_ascii=False))
